# Copyright 2026 Baumgärtner Marketing GmbH
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl).
"""Outputs of a statement (spec 7.5, 7.12, 9.7).

GT-09 (the fill-in aid renders with the notice, only filled fields and
accessible tables), GT-10 (workbook, evidence, list of the Kennzahlen, data
set and data access package Z3 are complete and show the amounts of the
statement; CSV with header line, semicolon, decimal comma, CRLF;
``index.xml`` well-formed with its mandatory elements, U-25), GT-11 (the
template of the procedure documentation holds the versions and the account
assignment), the registers of G-08 and the check protocol; GT-02 for the
outputs (no data of another business). SKR03 and SKR04; synthetic cases
(D6).
"""

import csv
import io
import json
import zipfile
from decimal import Decimal
from xml.etree import ElementTree

import openpyxl
from lxml import html

from odoo import Command
from odoo.exceptions import AccessError, UserError
from odoo.tests import new_test_user, tagged
from odoo.tools import mute_logger

from odoo.addons.account.tests.common import (
    AccountTestInvoicingCommon,
    AccountTestInvoicingHttpCommon,
)

from .scenario_common import (
    SKR03_ACCOUNTS,
    SKR03_TAXES,
    SKR04_ACCOUNTS,
    SKR04_TAXES,
)
from .statement_common import D, StatementCommon

KU = "EUER.20.111"
GIFT = "EUER.25.174"
ENTERTAINMENT = "EUER.25.175"
MEALS = "EUER.25.171"
LOW_VALUE = "EUER.25.132"
DEPRECIATION = "EUER.25.130"
WITHDRAWAL = "EUER.29.122"
PROFIT = "EUER.G.219"
INCOME = "EUER.20.159"


def _cells(table, tag):
    return [cell.text_content().strip() for cell in table.iter(tag)]


class OutputScenarios(StatementCommon):
    def scenario(self):
        """Small business 2026: revenue, expense, gifts to a recipient and
        give-aways, entertainment, a private withdrawal, a camera (AVEÜR), a
        low-value asset, an asset pool of 2025 and a profit correction with
        its own receipt."""
        sale = self.sale("2026-03-10", 1000)
        self.bank("2026-03-20", 1000, matches=[(sale, 1000)])
        bill = self.purchase("2026-04-10", 200)
        self.bank("2026-04-15", -200, matches=[(bill, 200)])
        recipient = self.env["res.partner"].create({"name": "Gift Recipient"})
        gift = self.invoice("in_invoice", "2026-05-02", [("gifts", 30)])
        gift.l10n_de_euer_gift_recipient_id = recipient
        self.bank("2026-05-05", -30, matches=[(gift, 30)])
        bulk = self.invoice("in_invoice", "2026-05-03", [("gifts", 3, (), 2)])
        bulk.l10n_de_euer_gift_bulk = True
        self.bank("2026-05-06", -6, matches=[(bulk, 6)])
        meal = self.invoice(
            "in_invoice",
            "2026-05-10",
            [("entertainment", 70), ("entertainment_nd", 30)],
        )
        self.bank("2026-05-12", -100, matches=[(meal, 100)])
        self.bank("2026-06-01", -100, extra=[("private_w", 100)])
        camera_bill = self.purchase("2026-07-15", 2150, account="asset")
        self.bank("2026-07-20", -2150, matches=[(camera_bill, 2150)])
        Asset = self.env["l10n_de_euer.asset"]
        camera = Asset.create(
            {
                "name": "Kamera",
                "code": "INV-1",
                "company_id": self.company.id,
                "delivery_date": "2026-07-15",
                "cost_net": 2150.0,
                "method": "linear",
                "useful_life_months": 84,
                "aveuer_group": "other_movable",
                "move_line_ids": [Command.set(camera_bill.invoice_line_ids.ids)],
            }
        )
        camera.action_activate()
        gwg = Asset.create(
            {
                "name": "Blitzgerät",
                "code": "INV-2",
                "company_id": self.company.id,
                "asset_kind": "gwg",
                "method": "none",
                "delivery_date": "2026-08-01",
                "cost_net": 500.0,
                "aveuer_group": "other_movable",
            }
        )
        gwg.action_activate()
        pooled = Asset.create(
            {
                "name": "Stativ",
                "company_id": self.company.id,
                "asset_kind": "pool",
                "method": "none",
                "delivery_date": "2025-06-01",
                "cost_net": 500.0,
                "aveuer_group": "other_movable",
            }
        )
        pooled.action_activate()
        statement = self.statement()
        receipt = self.env["ir.attachment"].create(
            {"name": "Eigenbeleg.txt", "raw": b"Reisekosten"}
        )
        correction = self.correction(
            statement,
            "meals",
            days_8h=1,
            attachment_ids=[Command.set(receipt.ids)],
        )
        statement.action_compute()
        return {
            "statement": statement,
            "sale": sale,
            "recipient": recipient,
            "camera": camera,
            "gwg": gwg,
            "pooled": pooled,
            "correction": correction,
            "receipt": receipt,
        }

    def render(self, xml_id, statement):
        content, kind = self.env["ir.actions.report"]._render_qweb_html(
            xml_id, statement.ids
        )
        self.assertEqual(kind, "html")
        return html.fromstring(content)

    def assertAccessibleTables(self, document):
        """O-11, 7.12: every table has a caption, column headers with scope
        col, a header cell with scope row per body row, numbers right
        aligned, no information only by colour."""
        tables = document.xpath("//table")
        self.assertTrue(tables)
        # no layout tables in the HTML view: only the tables of the output
        self.assertEqual(
            tables, document.xpath("//table[contains(@class, 'o_l10n_de_euer')]")
        )
        for table in tables:
            self.assertTrue(
                table.xpath("./caption") and table.xpath("./caption")[0].text_content(),
                html.tostring(table)[:300],
            )
            for header in table.xpath("./thead//th"):
                self.assertEqual(header.get("scope"), "col")
            for row in table.xpath("./tbody/tr"):
                self.assertTrue(row.xpath("./th[@scope='row']"))
        self.assertFalse(document.xpath("//*[contains(@class, 'text-danger')]"))

    def line(self, statement, code):
        return statement.line_ids.filtered(lambda line: line.code == code)

    # ------------------------------------------------------------ GT-09
    def test_gt09_fill_aid_html(self):
        """GT-09: the fill-in aid renders as HTML with the notice, the
        status, only the filled fields (in the order of the form) and the
        mandatory header fields, sums as control values, the previous year,
        the AVEÜR per group and per asset, in accessible tables."""
        data = self.scenario()
        statement = data["statement"]
        document = self.render("l10n_de_euer.report_fill_aid", statement)
        text = document.text_content()
        self.assertIn("no paper filing", text)
        self.assertIn("not finalised", text)
        self.assertIn("Mein ELSTER", text)
        self.assertAccessibleTables(document)
        # official labels are German and marked as such
        labels = document.xpath("//table[contains(@class, 'o_l10n_de_euer_fill_aid')]")
        self.assertTrue(labels)
        rows = {}
        for table in labels:
            for row in table.xpath("./tbody/tr"):
                cells = row.xpath("./td|./th")
                rows[cells[1].text_content().strip() or cells[2].text_content()] = row
                self.assertEqual(row.xpath("./th")[0].get("lang"), "de")
        # the filled fields with their amounts, right aligned
        ku_row = rows["111"]
        amount = ku_row.xpath("./td")[2]
        self.assertEqual(amount.text_content().strip(), "1.000,00")
        self.assertIn("text-end", amount.get("class"))
        self.assertIn("174", rows)
        self.assertIn("171", rows)
        self.assertIn("132", rows)
        # empty fields are left out (R-46): no VAT for a small business
        self.assertNotIn("140", rows)
        self.assertNotIn("185", rows)
        # mandatory header values and control values
        self.assertIn("Fotografie", text)
        self.assertIn("Control value", rows["159"].text_content())
        self.assertIn("Control value", rows["219"].text_content())
        # the order of the official form
        filled = statement._get_filled_lines()
        euer = [line.kz for line in filled if line.form_type == "euer" and line.kz]
        rendered = [
            row.xpath("./td")[1].text_content().strip()
            for row in labels[0].xpath("./tbody/tr")
            if row.xpath("./td")[1].text_content().strip()
        ]
        self.assertEqual(rendered, euer)
        # AVEÜR per group and per asset
        captions = [caption.text_content() for caption in document.xpath("//caption")]
        self.assertTrue(any("Anlage AVEÜR" in caption for caption in captions))
        assets = document.xpath("//table[contains(@class, 'o_l10n_de_euer_assets')]")
        self.assertIn("Kamera", assets[0].text_content())
        self.assertIn("153,57", assets[0].text_content())

    def test_gt09_fill_aid_minus_as_text_and_previous_year(self):
        """GT-09, 7.12: a loss is written with the minus in the text (never by
        colour), the previous year is shown, the PDF uses the same
        template; a statement in draft has no outputs."""
        sale = self.sale("2025-03-10", 500)
        self.bank("2025-03-20", 500, matches=[(sale, 500)])
        previous = self.statement(2025)
        previous.action_compute()
        sale = self.sale("2026-03-10", 100)
        self.bank("2026-03-20", 100, matches=[(sale, 100)])
        bill = self.purchase("2026-04-10", 300)
        self.bank("2026-04-15", -300, matches=[(bill, 300)])
        statement = self.statement()
        with self.assertRaises(UserError):
            statement.action_view_fill_aid()
        with self.assertRaises(UserError):
            statement._get_export("json")
        statement.action_compute()
        self.assertEqual(statement.profit, -200.0)
        data = statement._get_fill_aid_data()
        rows = {
            row["code"]: row for section in data["sections"] for row in section["rows"]
        }
        self.assertEqual(rows[PROFIT]["value"], "-200,00")
        self.assertTrue(rows[PROFIT]["negative"])
        self.assertEqual(rows[KU]["prev"], "500,00")
        document = self.render("l10n_de_euer.report_fill_aid", statement)
        self.assertIn("-200,00", document.text_content())
        self.assertAccessibleTables(document)
        content, kind = self.env["ir.actions.report"]._render_qweb_pdf(
            "l10n_de_euer.action_report_fill_aid", statement.ids
        )
        self.assertIn(b"-200,00", content)
        action = statement.action_view_fill_aid()
        self.assertEqual(action["report_type"], "qweb-html")
        printed = statement.with_context(discard_logo_check=True)
        self.assertEqual(printed.action_print_fill_aid()["report_type"], "qweb-pdf")

    def test_gt09_header_values_and_override(self):
        """GT-03 in the fill-in aid: a missing mandatory header value is
        listed and marked; key values show their meaning; an override
        shows the computed value and the reason."""
        self.company.l10n_de_euer_business_type = False
        sale = self.sale("2026-03-10", 1000)
        self.bank("2026-03-20", 1000, matches=[(sale, 1000)])
        statement = self.statement()
        self.answer(statement)
        statement.action_compute()
        line = self.line(statement, KU)
        line.write(
            {"source": "override", "override_amount": 1100, "override_reason": "Test"}
        )
        rows = {
            row["code"]: row
            for section in statement._get_fill_aid_data()["sections"]
            for row in section["rows"]
        }
        self.assertTrue(rows["EUER.15.100"]["missing"])
        self.assertEqual(rows["EUER.15.103"]["meaning"], "2 Trade or Business")
        self.assertEqual(rows["EUER.15.120"]["value"], "2")
        self.assertEqual(rows["EUER.15.120"]["meaning"], "No")
        self.assertEqual(rows[KU]["value"], "1.100,00")
        self.assertIn("1.000,00", rows[KU]["note"])
        self.assertIn("Test", rows[KU]["note"])
        document = self.render("l10n_de_euer.report_fill_aid", statement)
        self.assertIn("Missing.", document.text_content())

    # ------------------------------------------------------------ GT-10
    def test_gt10_evidence_matches_statement(self):
        """GT-10, G-03: the evidence lists every slice of each field with
        cash date, year, document, partner, account, share and source; its
        sum per field is the value of the statement."""
        data = self.scenario()
        statement = data["statement"]
        sections = {
            section["code"]: section
            for section in statement._get_evidence_data()["sections"]
        }
        for code, section in sections.items():
            line = self.line(statement, code)
            if line:
                self.assertEqual(section["total"], D(str(line.amount_computed)), code)
        ku = sections[KU]
        self.assertEqual(len(ku["rows"]), 1)
        row = ku["rows"][0]
        self.assertEqual(row["document"], data["sale"].name)
        self.assertEqual(row["date"].isoformat(), "2026-03-20")
        self.assertEqual(row["year"], 2026)
        self.assertEqual(row["amount"], D("1000.00"))
        self.assertEqual(row["share_text"], "100,00 %")
        self.assertTrue(row["account"])
        self.assertEqual(row["source"], "Document Paid")
        meals = sections[MEALS]["rows"][0]
        self.assertEqual(meals["source"], "Profit Correction")
        self.assertEqual(meals["adjustment_id"], data["correction"].id)
        # PDF (HTML in tests) and spreadsheet
        document = self.render("l10n_de_euer.report_evidence", statement)
        self.assertAccessibleTables(document)
        self.assertIn(data["sale"].name, document.text_content())
        name, mimetype, content = statement._get_export("evidence")
        self.assertTrue(name.endswith(".xlsx"))
        sheet = openpyxl.load_workbook(io.BytesIO(content)).active
        header = [cell.value for cell in sheet[1]]
        self.assertEqual(header[0], "Field")
        amounts = {}
        for values in sheet.iter_rows(min_row=2, values_only=True):
            amounts[values[0]] = amounts.get(values[0], 0) + values[11]
        self.assertAlmostEqual(amounts[KU], 1000.0)
        self.assertAlmostEqual(amounts[GIFT], 36.0)

    def test_gt10_workbook(self):
        """GT-10: the workbook holds the sheets of spec 7.5 with the values
        of the statement."""
        data = self.scenario()
        statement = data["statement"]
        name, mimetype, content = statement._get_export("workbook")
        self.assertEqual(
            mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("2026", name)
        workbook = openpyxl.load_workbook(io.BytesIO(content))
        self.assertEqual(
            workbook.sheetnames,
            [
                "Statement",
                "EÜR",
                "AVEÜR",
                "Evidence",
                "Asset Register",
                "Register of Low-Value Assets",
                "Asset Pools",
                "Withdrawals and Contributions",
                "Gifts per Recipient",
                "Business Entertainment",
                "Check Protocol",
                "Profit Corrections",
            ],
        )

        def rows(sheet):
            values = list(workbook[sheet].iter_rows(values_only=True))
            header = values[0]
            return [dict(zip(header, row, strict=True)) for row in values[1:]]

        euer = {row["Kz"]: row for row in rows("EÜR") if row["Kz"]}
        for line in statement._get_filled_lines().filtered(
            lambda line: line.form_type == "euer" and line.kz and line.has_value
        ):
            if line.field_type in ("text", "key", "date"):
                self.assertEqual(euer[line.kz]["Value"], line.value_text)
            else:
                self.assertAlmostEqual(euer[line.kz]["Amount"], line.amount)
        self.assertAlmostEqual(euer["111"]["Amount"], 1000.0)
        self.assertEqual(euer["111"]["Source"], "Computed")
        aveuer = rows("AVEÜR")
        self.assertTrue(aveuer)
        register = {row["Asset"]: row for row in rows("Asset Register")}
        self.assertAlmostEqual(register["Kamera"]["Depreciation"], 153.57)
        self.assertEqual(register["Kamera"]["Inventory Number"], "INV-1")
        self.assertEqual(
            [row["Asset"] for row in rows("Register of Low-Value Assets")],
            ["Blitzgerät"],
        )
        pools = rows("Asset Pools")
        self.assertEqual(pools[0]["Year of Formation"], 2025)
        self.assertAlmostEqual(pools[0]["Release of the Year"], 100.0)
        withdrawals = rows("Withdrawals and Contributions")
        # the private withdrawal and the contribution of the meal allowance
        self.assertEqual(
            [(row["Kind"], row["Amount"]) for row in withdrawals],
            [("Withdrawal", 100.0), ("Contribution", 14.0)],
        )
        gifts = rows("Gifts per Recipient")
        self.assertEqual(gifts[0]["Recipient"], "Gift Recipient")
        self.assertAlmostEqual(gifts[0]["Total of the Recipient"], 30.0)
        self.assertEqual(gifts[-1]["Give-Away"], "Yes")
        self.assertFalse(gifts[-1]["Recipient"])
        entertainment = rows("Business Entertainment")
        self.assertEqual(sorted(row["Amount"] for row in entertainment), [30.0, 70.0])
        corrections = rows("Profit Corrections")
        self.assertEqual(corrections[0]["Own Receipt"], "Eigenbeleg.txt")
        status = {
            row[0]: row[1] for row in workbook["Statement"].iter_rows(values_only=True)
        }
        self.assertEqual(status["State"], "Computed")

    def test_gt10_elster_csv(self):
        """GT-10: list of the Kennzahlen: UTF-8, semicolon, decimal comma,
        only filled fields, the columns of spec 7.5."""
        data = self.scenario()
        statement = data["statement"]
        name, mimetype, content = statement._get_export("elster")
        self.assertTrue(name.endswith(".csv"))
        self.assertTrue(mimetype.startswith("text/csv"))
        text = content.decode("utf-8-sig")
        self.assertIn("\r\n", text)
        rows = list(csv.reader(io.StringIO(text), delimiter=";"))
        self.assertEqual(
            rows[0],
            [
                "Year",
                "Form",
                "Form Version",
                "Box",
                "Line",
                "Kz",
                "Code",
                "Official Label",
                "Amount",
                "ELSTER Field ID",
                "XPath",
                "Verified",
                "Source",
            ],
        )
        by_code = {row[6]: row for row in rows[1:]}
        self.assertEqual(by_code[KU][8], "1000,00")
        self.assertEqual(by_code[KU][5], "111")
        self.assertEqual(by_code[KU][0], "2026")
        self.assertEqual(by_code[KU][11], "No")
        self.assertEqual(by_code["EUER.15.100"][8], "Fotografie")
        self.assertNotIn("EUER.20.140", by_code)
        filled = statement._get_filled_lines().filtered(lambda line: line.has_value)
        self.assertEqual(set(by_code), set(filled.mapped("code")))

    def test_gt10_json(self):
        """GT-10: data set ``l10n_de_euer/1``: year, form versions, business,
        fields, AVEÜR per group and per asset; the hash of the finalisation
        once finalised."""
        data = self.scenario()
        statement = data["statement"]
        name, mimetype, content = statement._get_export("json")
        self.assertEqual(mimetype, "application/json")
        dataset = json.loads(content)
        self.assertEqual(dataset["schema"], "l10n_de_euer/1")
        self.assertEqual(dataset["jahr"], 2026)
        self.assertIsNone(dataset["hash"])
        self.assertEqual(dataset["betrieb"]["art"], "Fotografie")
        self.assertEqual(
            {form["formular"] for form in dataset["formularversion"]},
            {"euer", "aveuer", "sz"},
        )
        fields = {row["code"]: row for row in dataset["felder"]}
        self.assertEqual(fields[KU]["betrag"], "1000.00")
        self.assertEqual(fields[KU]["kz"], "111")
        self.assertEqual(fields[KU]["quelle"], "berechnet")
        self.assertEqual(fields[PROFIT]["betrag"], f"{statement.profit:.2f}")
        self.assertTrue(fields[PROFIT]["kontrollwert"])
        self.assertEqual(fields["EUER.15.103"]["wert"], "2")
        groups = {row["code"]: row for row in dataset["aveuer"]["gruppen"]}
        self.assertTrue(groups)
        self.assertTrue(all(code.startswith("AVEUER.") for code in groups))
        details = dataset["aveuer"]["einzelangaben"]
        self.assertEqual(details[0]["bezeichnung"], "Kamera")
        self.assertEqual(details[0]["afa"], "153.57")
        self.assertEqual(dataset["kontrollwerte"]["gewinn"], f"{statement.profit:.2f}")

    def test_gt10_z3(self):
        """GT-10, G-09: data access package: CSV with header line, semicolon,
        decimal comma and CRLF for slices, field values, account assignment
        and the registers; ``index.xml`` describes every file with its
        columns; the description of the fields as JSON; the slices add up
        to the values of the statement."""
        data = self.scenario()
        statement = data["statement"]
        name, mimetype, content = statement._get_export("z3")
        self.assertEqual(mimetype, "application/zip")
        archive = zipfile.ZipFile(io.BytesIO(content))
        names = archive.namelist()
        for file in (
            "slices.csv",
            "statement_fields.csv",
            "account_assignment.csv",
            "asset_register.csv",
            "low_value_assets.csv",
            "asset_pools.csv",
            "withdrawals_contributions.csv",
            "gifts.csv",
            "entertainment.csv",
            "profit_corrections.csv",
            "checks.csv",
            "index.xml",
            "field_description.json",
        ):
            self.assertIn(file, names)
        index = ElementTree.fromstring(archive.read("index.xml"))
        tables = {table.findtext("URL"): table for table in index.iter("Table")}
        self.assertEqual(set(tables), {name for name in names if name.endswith(".csv")})
        for file, table in tables.items():
            raw = archive.read(file).decode("utf-8")
            self.assertTrue(raw.endswith("\r\n"))
            self.assertNotIn("\r\n\r\n", raw)
            header = raw.split("\r\n", 1)[0].split(";")
            described = [
                element.findtext("Name")
                for element in table.find("VariableLength")
                if element.tag in ("VariablePrimaryKey", "VariableColumn")
            ]
            self.assertEqual(header, described, file)
            self.assertEqual(table.findtext("Validity/Range/From"), "01.01.2026")
        slices = list(
            csv.DictReader(
                io.StringIO(archive.read("slices.csv").decode()), delimiter=";"
            )
        )
        self.assertEqual(len(slices), len(statement.run_id.slice_ids))
        total = sum(
            Decimal(row["amount"].replace(",", "."))
            for row in slices
            if row["field_code"] == KU and row["attribution_year"] == "2026"
        )
        self.assertEqual(total, Decimal("1000.00"))
        self.assertTrue(any("," in row["amount"] for row in slices))
        fields = list(
            csv.DictReader(
                io.StringIO(archive.read("statement_fields.csv").decode()),
                delimiter=";",
            )
        )
        values = {row["code"]: row for row in fields}
        self.assertEqual(values[KU]["amount"], "1000,00")
        assignment = archive.read("account_assignment.csv").decode()
        revenue = self.acc("ku_revenue")
        self.assertIn(f";{revenue.code};", assignment)
        description = json.loads(archive.read("field_description.json"))
        self.assertEqual(
            {table["file"] for table in description["tables"]}, set(tables)
        )
        gifts = archive.read("gifts.csv").decode()
        self.assertIn("Gift Recipient", gifts)

    def test_gt10_check_protocol(self):
        """Check protocol (spec 7.5): all checks with their state and the
        acknowledgement (reason, user, time), as PDF and XLSX."""
        data = self.scenario()
        statement = data["statement"]
        self.answer(statement)
        statement.action_check()
        warning = statement.check_ids.filtered(
            lambda check: check.severity == "warning"
        )[:1]
        self.assertTrue(warning)
        warning.acknowledge_reason = "Checked by hand"
        warning.action_acknowledge()
        protocol = statement._get_check_data()
        self.assertEqual(len(protocol["rows"]), len(statement.check_ids))
        row = next(row for row in protocol["rows"] if row["id"] == warning.id)
        self.assertEqual(row["status"], "Acknowledged")
        self.assertEqual(row["reason"], "Checked by hand")
        self.assertEqual(row["acknowledged_by"], self.env.user.name)
        document = self.render("l10n_de_euer.report_check_protocol", statement)
        self.assertAccessibleTables(document)
        self.assertIn("Checked by hand", document.text_content())
        name, mimetype, content = statement._get_export("checks")
        sheet = openpyxl.load_workbook(io.BytesIO(content)).active
        codes = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sorted(codes), sorted(statement.check_ids.mapped("code")))

    def test_gt10_finalised_outputs(self):
        """After the finalisation the outputs show the state, the hash and
        the versions of the finalisation."""
        sale = self.sale("2026-03-10", 1000)
        self.bank("2026-03-20", 1000, matches=[(sale, 1000)])
        statement = self.finalize(self.statement())
        dataset = json.loads(statement._get_export("json")[2])
        self.assertEqual(dataset["hash"], statement.payload_sha256)
        self.assertEqual(dataset["status"]["kettenhash"], statement.chain_hash)
        self.assertEqual(dataset["status"]["modulversion"], statement.module_version)
        document = self.render("l10n_de_euer.report_fill_aid", statement)
        text = document.text_content()
        self.assertIn(statement.payload_sha256, text)
        self.assertIn(statement.chain_hash, text)
        self.assertNotIn("not finalised", text)
        archive = zipfile.ZipFile(io.BytesIO(statement._get_export("z3")[2]))
        comment = ElementTree.fromstring(archive.read("index.xml")).findtext(
            "DataSupplier/Comment"
        )
        self.assertIn("Finalised", comment)
        # the outputs never change finalised data
        self.assertEqual(statement._verify_integrity(), [])
        for kind in ("workbook", "evidence", "checks", "elster", "json", "z3"):
            action = getattr(
                statement,
                {
                    "workbook": "action_export_workbook",
                    "evidence": "action_export_evidence",
                    "checks": "action_export_checks",
                    "elster": "action_export_elster",
                    "json": "action_export_json",
                    "z3": "action_export_z3",
                }[kind],
            )()
            self.assertEqual(
                action["url"], f"/l10n_de_euer/statement/{statement.id}/{kind}"
            )
        with self.assertRaises(UserError):
            statement._get_export("pdf")

    def test_sheet_names(self):
        """Sheet names of the workbooks: no characters Excel refuses, at most
        31 characters, unique."""
        used = set()
        self.assertEqual(self.Statement._sheet_name("A/B: [x]?", used), "A B   x")
        self.assertEqual(self.Statement._sheet_name("a/b: [x]?", used), "a b   x 2")
        long_name = self.Statement._sheet_name("x" * 40, used)
        self.assertEqual(len(long_name), 31)
        self.assertEqual(len(self.Statement._sheet_name("x" * 40, used)), 31)
        self.assertEqual(self.Statement._sheet_name("", used), "Sheet")

    # ------------------------------------------------------------ GT-11
    def test_gt11_procedure_documentation(self):
        """GT-11: the template of the procedure documentation holds data
        flow, calculation rules with the parameters of the year, the account
        assignment, settings, roles, the lock concept with the limit of
        database administrators and the versions."""
        data = self.scenario()
        statement = data["statement"]
        procedure = statement._get_procedure_data()
        revenue = self.acc("ku_revenue")
        codes = {row["account_code"]: row for row in procedure["assignment"]}
        self.assertIn(revenue.code, codes)
        self.assertEqual(codes[revenue.code]["category_code"], "REV_KU")
        self.assertIn("Kz 111", codes[revenue.code]["fields"])
        self.assertEqual(procedure["status"]["module_version"], "19.0.1.0.0")
        self.assertEqual(
            procedure["status"]["ruleset_version"], statement._get_ruleset_version()
        )
        self.assertTrue(procedure["parameters"])
        self.assertTrue(procedure["roles"])
        self.assertIn(
            self.env.user.name,
            " ".join(row["users"] for row in procedure["roles"]),
        )
        document = self.render("l10n_de_euer.report_procedure", statement)
        text = document.text_content()
        for part in (
            "Data flow",
            "Calculation rules",
            "Account assignment",
            "Settings",
            "Roles",
            "Lock concept",
            "database administrators",
            "Versions",
            "19.0.1.0.0",
            statement._get_ruleset_version(),
            revenue.code,
        ):
            self.assertIn(part, text)
        for form in statement.form_ids:
            self.assertIn(form.name, text)
        self.assertAccessibleTables(document)
        self.assertEqual(statement.action_view_procedure()["report_type"], "qweb-html")
        printed = statement.with_context(discard_logo_check=True)
        for action in (
            printed.action_print_procedure(),
            printed.action_print_evidence(),
            printed.action_print_checks(),
        ):
            self.assertEqual(action["report_type"], "qweb-pdf")

    # ------------------------------------------------------------ GT-02
    def test_gt02_outputs_of_one_business(self):
        """GT-02: the outputs of a statement hold only the data of its
        business; a user of another business cannot read them."""
        data = self.scenario()
        statement = data["statement"]
        other = self.setup_other_company(name="Other Business")
        company_b = other["company"]
        company_b.l10n_de_euer_enabled = True
        self.env["l10n_de_euer.vat.period"].create(
            {"company_id": company_b.id, "date_from": "2020-01-01", "regime": "ku"}
        )
        secret = self.env["res.partner"].create({"name": "Secret Customer"})
        revenue_b = self.env.ref(
            f"account.{company_b.id}_{self.ACCOUNTS['ku_revenue']}"
        )
        self.bank(
            "2026-05-05",
            70,
            extra=[(revenue_b, -70)],
            journal=other["default_journal_bank"],
            partner=secret,
            ref="Secret payment",
        )
        statement_b = self.statement(company=company_b)
        statement_b.action_compute()
        self.assertEqual(statement_b.total_income, 70.0)
        statement.action_compute()
        for kind in ("workbook", "elster", "json", "z3", "evidence", "checks"):
            content = statement._get_export(kind)[2]
            if kind in ("workbook", "evidence", "checks"):
                workbook = openpyxl.load_workbook(io.BytesIO(content))
                text = " ".join(
                    str(value)
                    for sheet in workbook
                    for row in sheet.iter_rows(values_only=True)
                    for value in row
                    if value
                )
            elif kind == "z3":
                archive = zipfile.ZipFile(io.BytesIO(content))
                text = " ".join(
                    archive.read(name).decode() for name in archive.namelist()
                )
            else:
                text = content.decode("utf-8-sig")
            self.assertNotIn("Secret", text, kind)
            self.assertNotIn("Other Business", text, kind)
        for report in (
            "l10n_de_euer.report_fill_aid",
            "l10n_de_euer.report_evidence",
            "l10n_de_euer.report_procedure",
        ):
            text = self.render(report, statement).text_content()
            self.assertNotIn("Secret", text, report)
        outsider = new_test_user(
            self.env,
            login="output_outsider",
            groups="account.group_account_readonly",
            company_id=company_b.id,
            company_ids=[Command.set(company_b.ids)],
        )
        with self.assertRaises(AccessError):
            statement.with_user(outsider)._get_export("json")
        own = statement_b.with_user(outsider)._get_export("z3")[2]
        slices_b = zipfile.ZipFile(io.BytesIO(own)).read("slices.csv").decode()
        self.assertIn("Secret Customer", slices_b)
        self.assertNotIn(data["sale"].name, slices_b)


@tagged("post_install_l10n", "post_install", "-at_install")
class TestStatementOutputSkr03(OutputScenarios, AccountTestInvoicingCommon):
    allow_inherited_tests_method = True
    CHART = "de_skr03"
    ACCOUNTS = SKR03_ACCOUNTS
    TAXES = SKR03_TAXES

    @classmethod
    @AccountTestInvoicingCommon.setup_country("de")
    @AccountTestInvoicingCommon.setup_chart_template("de_skr03")
    def setUpClass(cls):
        super().setUpClass()
        cls.setup_finalisation()


@tagged("post_install_l10n", "post_install", "-at_install")
class TestStatementOutputSkr04(OutputScenarios, AccountTestInvoicingCommon):
    allow_inherited_tests_method = True
    CHART = "de_skr04"
    ACCOUNTS = SKR04_ACCOUNTS
    TAXES = SKR04_TAXES

    @classmethod
    @AccountTestInvoicingCommon.setup_country("de")
    @AccountTestInvoicingCommon.setup_chart_template("de_skr04")
    def setUpClass(cls):
        super().setUpClass()
        cls.setup_finalisation()


@tagged("post_install_l10n", "post_install", "-at_install")
class TestStatementOutputDownload(StatementCommon, AccountTestInvoicingHttpCommon):
    """The downloads through the controller (spec 7.2 controllers): the
    file with its name, an auditor may download, a user of another
    business may not, an unknown kind is not found."""

    CHART = "de_skr03"
    ACCOUNTS = SKR03_ACCOUNTS
    TAXES = SKR03_TAXES

    @classmethod
    @AccountTestInvoicingCommon.setup_country("de")
    @AccountTestInvoicingCommon.setup_chart_template("de_skr03")
    def setUpClass(cls):
        super().setUpClass()
        cls.setup_finalisation()
        cls.auditor = new_test_user(
            cls.env,
            login="output_auditor",
            password="output_auditor",
            groups="account.group_account_readonly",
            company_id=cls.company.id,
            company_ids=[Command.set(cls.company.ids)],
        )
        other = cls.env["res.company"].create({"name": "Other Download Business"})
        cls.outsider = new_test_user(
            cls.env,
            login="output_other",
            password="output_other",
            groups="account.group_account_readonly",
            company_id=other.id,
            company_ids=[Command.set(other.ids)],
        )

    def test_download(self):
        sale = self.sale("2026-03-10", 1000)
        self.bank("2026-03-20", 1000, matches=[(sale, 1000)])
        statement = self.statement()
        statement.action_compute()
        url = f"/l10n_de_euer/statement/{statement.id}"
        self.authenticate("output_auditor", "output_auditor")
        response = self.url_open(f"{url}/elster")
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response.headers["Content-Disposition"])
        self.assertIn(".csv", response.headers["Content-Disposition"])
        self.assertIn("1000,00", response.content.decode("utf-8-sig"))
        response = self.url_open(f"{url}/z3")
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "index.xml", zipfile.ZipFile(io.BytesIO(response.content)).namelist()
        )
        self.assertEqual(self.url_open(f"{url}/unknown").status_code, 404)
        self.assertEqual(
            self.url_open("/l10n_de_euer/statement/999999/json").status_code, 404
        )
        self.authenticate("output_other", "output_other")
        with mute_logger("odoo.http"):
            response = self.url_open(f"{url}/json")
        self.assertNotEqual(response.status_code, 200)
        self.assertNotIn(b"l10n_de_euer/1", response.content)
